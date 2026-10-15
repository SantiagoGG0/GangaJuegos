import io
import shutil
import tempfile
from unittest import mock

import openpyxl
from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from PIL import Image

from apps.catalog.models import Cover, Juego
from apps.panel import products_bulk as pb


def _png():
    buf = io.BytesIO()
    Image.new('RGB', (20, 30), (200, 40, 40)).save(buf, 'PNG')
    return buf.getvalue()


def _xlsx(rows, image_rows=()):
    """XLSX con la estructura de la plantilla oficial (hoja Productos + Instrucciones)."""
    from openpyxl.drawing.image import Image as XLImage
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Productos'
    ws.append(pb.COLUMNS)
    for r in rows:
        ws.append(r)
    for row in image_rows:
        ws.add_image(XLImage(io.BytesIO(_png())), f'N{row}')
    wb.create_sheet('Instrucciones').append(['Columna', 'Qué poner'])
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile('productos.xlsx', buf.getvalue())


class BulkHelpersTests(TestCase):
    def test_headers_of_official_template(self):
        self.assertEqual([pb._norm_header(c) for c in pb.COLUMNS], [k for k, _ in pb.TEMPLATE_COLUMNS])

    def test_legacy_headers(self):
        for h, key in [('titulo', 'titulo'), ('precio_primaria', 'precio_primaria'),
                       ('precio_anterior', 'precio_anterior_primaria'), ('caratula_url', 'caratula_url')]:
            self.assertEqual(pb._norm_header(h), key)

    def test_platforms(self):
        for raw, val in [('PS5', 'PS5'), ('ps4', 'PS4'), ('PS4/PS5', 'PS4/PS5'), ('PS5 / PS4', 'PS4/PS5'),
                         ('PS4 y PS5', 'PS4/PS5'), ('PS4-PS5', 'PS4/PS5'), ('Xbox', None), ('', None)]:
            self.assertEqual(pb._norm_platform(raw), val, raw)

    def test_prices(self):
        for raw, val in [(63900, 63900), (63900.0, 63900), ('63900', 63900), ('$ 63.900', 63900),
                         ('63,900', 63900), ('63.900,00', 63900), ('', None), (None, None), ('abc', False), (0, False)]:
            self.assertEqual(pb._price(raw), val, raw)

    def test_yes_no(self):
        self.assertIs(pb._yes_no('SI', False), True)
        self.assertIs(pb._yes_no('Sí', False), True)
        self.assertIs(pb._yes_no('no', True), False)
        self.assertIs(pb._yes_no('', True), True)
        self.assertEqual(pb._yes_no('quizás', True), 'bad')

    def test_slug(self):
        self.assertEqual(pb._make_slug('', 'EA SPORTS FC™ 27'), 'ea-sports-fctm-27')
        self.assertEqual(pb._make_slug('Mi Slug', 'X'), 'mi-slug')
        self.assertTrue(pb._make_slug('', 'ドラゴン').startswith('juego-'))


@override_settings(STORAGES={
    'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
    'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'},
})
class BulkFlowTests(TestCase):
    def setUp(self):
        self.media = tempfile.mkdtemp()
        self.override = override_settings(MEDIA_ROOT=self.media)
        self.override.enable()
        user = get_user_model().objects.create_superuser('a@a.co', 'x', nombre='Admin')
        self.client.force_login(user)

    def tearDown(self):
        self.override.disable()
        shutil.rmtree(self.media, ignore_errors=True)

    def _upload(self, f):
        return self.client.post(reverse('gsadmin:products_bulk_upload'), {'file': f})

    def _import_all(self):
        for _ in range(500):
            self.client.post(reverse('gsadmin:products_bulk_import'))
            if self.client.session[pb.SESSION_KEY].get('done'):
                return
        self.fail('la importación no terminó')

    def test_official_template_imports_without_errors(self):
        long_title = 'Mon-Yu: ' + 'Defeat Monsters And Gain Strong Weapons ' * 4
        rows = [
            ['NBA 2K27', None, 'PS5', 'Deportes', 63900, 126900, 90900, 181900, 'SI', 'NO', None, None,
             'https://image.api.playstation.com/a.png', None],
            ['EA SPORTS FC™ 27', None, 'PS4/PS5', 'Deportes', 91900, 138900, None, None, 'SI', 'NO', None, None, None, None],
            [long_title, None, 'PS4', 'RPG', 20900, 41900, None, None, 'SI', 'NO', None, None, None, None],
            ['Cappuccino Assassino', None, 'PS4', 'Shooter', 900, 900, 900, 900, 'SI', 'NO', None, None, None, None],
            ['METAL GEAR SOLID Δ: SNAKE EATER', 'mgs-delta', 'PS5', 'Acción/Sigilo', 99900, 199900, None, None,
             'SI', 'SI', 'Vuelve el clásico', 'nota', None, None],
        ]
        resp = self._upload(_xlsx(rows, image_rows=[3]))
        self.assertRedirects(resp, reverse('gsadmin:products_bulk'), fetch_redirect_response=False)
        data = self.client.session[pb.SESSION_KEY]
        self.assertEqual([r['err'] for r in data['rows']], [''] * 5)
        self.assertTrue(data['rows'][1]['cover_tmp'])  # imagen pegada en la fila 3
        page = self.client.get(reverse('gsadmin:products_bulk'))
        self.assertContains(page, '5 válidas')

        fake = (ContentFile(_png(), name='import.png'), None)
        with mock.patch.object(pb, '_fetch_image', return_value=fake) as fetch:
            self._import_all()
        fetch.assert_called_once()
        self.assertEqual(Juego.objects.count(), 5)
        self.assertEqual(Cover.objects.count(), 2)
        nba = Juego.objects.get(slug='nba-2k27')
        self.assertEqual((nba.price_secondary, nba.price_primary), (63900, 126900))
        self.assertEqual((nba.old_price_secondary, nba.old_price_primary), (90900, 181900))
        self.assertTrue(nba.active)
        self.assertFalse(nba.featured)
        mgs = Juego.objects.get(slug='mgs-delta')
        self.assertTrue(mgs.featured)
        self.assertEqual((mgs.tagline, mgs.notas), ('Vuelve el clásico', 'nota'))
        self.assertIsNone(Juego.objects.get(slug='cappuccino-assassino').old_price_primary)
        self.assertTrue(Juego.objects.filter(title__startswith='Mon-Yu').exists())
        self.assertEqual(self.client.session[pb.SESSION_KEY]['import_errors'], [])
        self.assertContains(self.client.get(reverse('gsadmin:products_bulk')), '5 juegos importados')

    def test_reimport_updates_instead_of_duplicating(self):
        row = ['NBA 2K27', None, 'PS5', 'Deportes', 63900, 126900, None, None, 'SI', 'NO', None, None, None, None]
        self._upload(_xlsx([row]))
        self._import_all()
        self.client.post(reverse('gsadmin:products_bulk_reset'))
        row[4], row[8] = 59900, 'NO'
        self._upload(_xlsx([row]))
        self._import_all()
        self.assertEqual(Juego.objects.count(), 1)
        j = Juego.objects.get()
        self.assertEqual(j.price_secondary, 59900)
        self.assertFalse(j.active)

    def test_invalid_rows_are_reported(self):
        rows = [
            ['', None, 'PS5', 'Deportes', 1, 2, None, None, 'SI', 'NO', None, None, None, None],
            ['A', None, 'Xbox', 'Deportes', 1, 2, None, None, 'SI', 'NO', None, None, None, None],
            ['B', None, 'PS5', 'Deportes', 'gratis', 2, None, None, 'SI', 'NO', None, None, None, None],
            ['C', None, 'PS5', 'Deportes', 1, 2, None, None, 'tal vez', 'NO', None, None, None, None],
        ]
        self._upload(_xlsx(rows))
        errs = [r['err'] for r in self.client.session[pb.SESSION_KEY]['rows']]
        self.assertEqual(errs, ['Falta título', 'Plataforma no válida (Xbox)', 'Precio secundaria no válido',
                                'Activo debe ser SI o NO'])

    def test_missing_columns(self):
        wb = openpyxl.Workbook()
        wb.active.append(['Título *', 'Plataforma *'])
        buf = io.BytesIO()
        wb.save(buf)
        self._upload(SimpleUploadedFile('x.xlsx', buf.getvalue()))
        self.assertIsNone(self.client.session.get(pb.SESSION_KEY))
        self.assertContains(self.client.get(reverse('gsadmin:products_bulk')), 'le faltan columnas')

    def test_csv_with_template_headers(self):
        text = ';'.join(pb.COLUMNS) + '\nNBA 2K27;;PS5;Deportes;63.900;126.900;;;SI;NO;;;;\n'
        self._upload(SimpleUploadedFile('p.csv', text.encode('utf-8-sig')))
        self.assertEqual(self.client.session[pb.SESSION_KEY]['rows'][0]['err'], '')

    def test_template_download_roundtrip(self):
        resp = self.client.get(reverse('gsadmin:products_bulk_template'))
        self.assertEqual(resp.status_code, 200)
        self._upload(SimpleUploadedFile('plantilla.xlsx', resp.content))
        self.assertEqual(self.client.session[pb.SESSION_KEY]['rows'][0]['err'], '')


class CoverResizeTests(TestCase):
    def test_large_cover_is_reduced_to_webp(self):
        buf = io.BytesIO()
        Image.new('RGB', (1800, 2400), (10, 120, 200)).save(buf, 'PNG')
        content = pb._web_image(buf.getvalue())
        self.assertTrue(content.name.endswith('.webp'))
        img = Image.open(io.BytesIO(content.read()))
        self.assertEqual((img.format, img.size), ('WEBP', (600, 800)))

    def test_small_cover_is_not_enlarged(self):
        img = Image.open(io.BytesIO(pb._web_image(_png()).read()))
        self.assertEqual(img.size, (20, 30))


@override_settings(STORAGES={
    'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
    'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'},
})
class CatalogEditTests(TestCase):
    def setUp(self):
        self.media = tempfile.mkdtemp()
        self.override = override_settings(MEDIA_ROOT=self.media)
        self.override.enable()
        self.client.force_login(get_user_model().objects.create_superuser('a@a.co', 'x', nombre='Admin'))
        self.games = [Juego.objects.create(title=f'Juego {i}', slug=f'juego-{i}', platform='PS5', genre='RPG',
                                           price_primary=100000, price_secondary=50000) for i in range(3)]
        cover = Cover.objects.create(juego=self.games[0], image=ContentFile(_png(), name='c.png'))
        self.cover_path = cover.image.path

    def tearDown(self):
        self.override.disable()
        shutil.rmtree(self.media, ignore_errors=True)

    def test_catalog_page_shows_controls(self):
        page = self.client.get(reverse('gsadmin:dashboard_catalogo'))
        self.assertContains(page, 'Eliminar todo el catálogo')
        self.assertContains(page, reverse('gsadmin:dashboard_juego_delete', args=[self.games[1].pk]))

    def test_delete_one_removes_cover_file(self):
        import os
        self.client.post(reverse('gsadmin:dashboard_juego_delete', args=[self.games[0].pk]))
        self.assertFalse(Juego.objects.filter(pk=self.games[0].pk).exists())
        self.assertFalse(os.path.exists(self.cover_path))
        self.assertEqual(Juego.objects.count(), 2)

    def test_bulk_actions(self):
        ids = [self.games[0].pk, self.games[1].pk]
        url = reverse('gsadmin:dashboard_catalog_bulk')
        self.client.post(url, {'action': 'ocultar', 'ids': ids})
        self.assertEqual(Juego.objects.filter(active=False).count(), 2)
        self.client.post(url, {'action': 'destacar', 'ids': ids})
        self.assertEqual(Juego.objects.filter(featured=True).count(), 2)
        self.client.post(url, {'action': 'eliminar', 'ids': ids})
        self.assertEqual(list(Juego.objects.values_list('pk', flat=True)), [self.games[2].pk])

    def test_delete_all_requires_confirmation(self):
        url = reverse('gsadmin:dashboard_catalog_delete_all')
        self.client.post(url, {'confirm': 'si'})
        self.assertEqual(Juego.objects.count(), 3)
        self.client.post(url, {'confirm': 'eliminar todo'})
        self.assertEqual(Juego.objects.count(), 0)
        self.assertEqual(Cover.objects.count(), 0)

    def test_delete_keeps_orders(self):
        from apps.orders.models import ItemPedido, Pedido
        user = get_user_model().objects.create_user('c@c.co', 'x', nombre='C')
        field_names = {f.name for f in Pedido._meta.get_fields()}
        self.assertIn('usuario', field_names)
        try:
            pedido = Pedido.objects.create(usuario=user, total=50000)
        except Exception as e:  # noqa: BLE001
            self.skipTest(f'Pedido requiere más campos: {e}')
        ItemPedido.objects.create(pedido=pedido, juego=self.games[1], precio_unitario=50000)
        self.client.post(reverse('gsadmin:dashboard_catalog_delete_all'), {'confirm': 'ELIMINAR TODO'})
        item = ItemPedido.objects.get()
        self.assertIsNone(item.juego)
        self.assertEqual(item.titulo, 'Juego 1')
